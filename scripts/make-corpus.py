#!/usr/bin/env python3
"""Erzeugt den Test-Korpus schwieriger PDFs (tests/fixtures). Benötigt: pip install pymupdf openpyxl python-docx"""
import os, io, pymupdf

OUT = os.path.join(os.path.dirname(__file__), '..', 'tests', 'fixtures')
os.makedirs(OUT, exist_ok=True)
FONT = os.path.join(os.path.dirname(__file__), '..', 'node_modules', 'pdfjs-dist', 'standard_fonts', 'LiberationSans-Regular.ttf')
P = lambda n: os.path.join(OUT, n)
TEXT = ("Werkbank ist eine PDF-Anwendung. Dieser Absatz dient als Testtext für Suche, Markierung "
        "und Schwärzung. Kontakt: max.mustermann@example.com, IBAN DE89 3704 0044 0532 0130 00. ")

def text_pages(doc, n, w=595, h=842, title='Kapitel'):
    for i in range(n):
        p = doc.new_page(width=w, height=h)
        p.insert_text((72, 90), f'{title} {i + 1}: Testdokument', fontsize=22)
        p.insert_textbox(pymupdf.Rect(72, 120, w - 72, 400), TEXT * 3, fontsize=11)
        p.draw_rect(pymupdf.Rect(72, 420, w - 72, 560), color=(0.2, 0.4, 0.8), fill=(0.9, 0.93, 1))
        p.insert_text((90, 450), f'Tabelle {i + 1}', fontsize=12)
    return doc

def save(doc, name, **kw):
    doc.save(P(name), garbage=3, deflate=True, **kw); doc.close()

# 1 Normaler Text
save(text_pages(pymupdf.open(), 6), 'text.pdf')

# 2 Gedrehte Seiten
d = text_pages(pymupdf.open(), 4)
for i, r in enumerate([0, 90, 180, 270]): d[i].set_rotation(r)
save(d, 'rotated.pdf')

# 3 Gemischte Formate
d = pymupdf.open()
for (w, h, t) in [(595, 842, 'A4'), (420, 595, 'A5'), (612, 792, 'Letter'), (1191, 842, 'A3 quer')]:
    p = d.new_page(width=w, height=h); p.insert_text((50, 80), f'Format {t}', fontsize=20)
save(d, 'mixed-sizes.pdf')

# 4 Verschlüsselt (Passwort zum Öffnen: geheim)
save(text_pages(pymupdf.open(), 2), 'encrypted.pdf', encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw='geheim', owner_pw='besitzer')

# 5 Nur Rechte eingeschränkt (öffnet ohne Passwort, Drucken/Kopieren verboten)
save(text_pages(pymupdf.open(), 2), 'restricted.pdf', encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='besitzer', user_pw='',
     permissions=pymupdf.PDF_PERM_ACCESSIBILITY)

# 6 Formular
d = pymupdf.open(); p = d.new_page()
p.insert_text((72, 80), 'Antrag auf Holzeinschlag', fontsize=18)
for i, (name, typ) in enumerate([('Name', pymupdf.PDF_WIDGET_TYPE_TEXT), ('Revier', pymupdf.PDF_WIDGET_TYPE_TEXT), ('Eilig', pymupdf.PDF_WIDGET_TYPE_CHECKBOX), ('Baumart', pymupdf.PDF_WIDGET_TYPE_COMBOBOX)]):
    w = pymupdf.Widget(); w.field_name = name; w.field_type = typ
    w.rect = pymupdf.Rect(180, 120 + i * 40, 400 if typ != pymupdf.PDF_WIDGET_TYPE_CHECKBOX else 196, 140 + i * 40)
    if typ == pymupdf.PDF_WIDGET_TYPE_COMBOBOX: w.choice_values = ['Fichte', 'Buche', 'Eiche']; w.field_value = 'Fichte'
    p.insert_text((72, 135 + i * 40), name, fontsize=11); p.add_widget(w)
save(d, 'form.pdf')

# 7 Vorhandene Anmerkungen
d = text_pages(pymupdf.open(), 1); p = d[0]
p.add_highlight_annot(p.search_for('Testtext')[0]).set_info(title='Kollegin')
a = p.add_text_annot((500, 100), 'Bitte prüfen'); a.set_info(title='Kollege'); a.update()
p.add_ink_annot([[(100, 600), (150, 650), (200, 600)]])
save(d, 'annotated.pdf')

# 8 Scan (nur Bild, klein gehalten)
src = text_pages(pymupdf.open(), 2); d = pymupdf.open()
for pg in src:
    pix = pg.get_pixmap(dpi=150, colorspace=pymupdf.csGRAY)
    np_ = d.new_page(width=595, height=842); np_.insert_image(np_.rect, stream=pix.tobytes('jpeg', jpg_quality=60))
save(d, 'scan.pdf')

# 9 Beschädigt (Querverweistabelle zerstört)
b = open(P('text.pdf'), 'rb').read()
i = b.rfind(b'startxref'); b = b[:i] + b'startxref\n999999\n%%EOF\n'
b = b.replace(b'/Root', b'/Root', 1)
open(P('damaged.pdf'), 'wb').write(b)

# 10 Lang
d = pymupdf.open()
for i in range(120):
    p = d.new_page(); p.insert_text((72, 90), f'Seite {i + 1} von 120', fontsize=16)
save(d, 'long.pdf')

# 11 Lesezeichen
d = text_pages(pymupdf.open(), 5)
d.set_toc([[1, 'Einleitung', 1], [1, 'Methoden', 2], [2, 'Inventur', 3], [1, 'Ergebnisse', 4], [1, 'Anhang', 5]])
save(d, 'outline.pdf')

# 12 Verschobene Seitenbox (MediaBox beginnt nicht bei 0,0 – klassische Geometrie-Falle)
d = pymupdf.open(); p = d.new_page(width=595, height=842)
p.insert_text((72, 90), 'Versetzte Seitenbox', fontsize=20)
p.insert_textbox(pymupdf.Rect(72, 120, 520, 300), TEXT, fontsize=11)
d.save(P('_tmp.pdf')); d.close()
d = pymupdf.open(P('_tmp.pdf'))
x = d.page_xref(0); d.xref_set_key(x, 'MediaBox', '[100 200 695 1042]'); d.xref_set_key(x, 'CropBox', '[120 220 675 1022]')
import re
refs = [int(n) for n in re.findall(r'(\d+) 0 R', d.xref_get_key(x, 'Contents')[1])]
d.update_stream(refs[0], b'q 1 0 0 1 100 200 cm\n' + d.xref_stream(refs[0]))
d.update_stream(refs[-1], d.xref_stream(refs[-1]) + b'\nQ')
save(d, 'offset-box.pdf'); os.remove(P('_tmp.pdf'))

# 13 Sonderzeichen mit eingebetteter Schrift
d = pymupdf.open(); p = d.new_page()
p.insert_text((72, 90), 'Grüße aus Łódź, Šumava und Ősz – 25 € · „Zitat“', fontsize=14, fontname='Lib', fontfile=FONT)
save(d, 'unicode.pdf')

# Office-Dateien und Bild
from docx import Document
doc = Document(); doc.add_heading('Testbericht Werkbank', 1); doc.add_paragraph('Fett und normaler Text mit Umlauten äöü ß.'); doc.save(P('brief.docx'))
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.append(['Baumart', 'Fläche ha', 'Vorrat m³']); ws.append(['Fichte', 12.5, 3400]); ws.append(['Buche', 8.0, 2100]); wb.save(P('bestand.xlsx'))
pymupdf.open(P('text.pdf'))[0].get_pixmap(dpi=60).save(P('foto.jpg'))
# Seite mit Logo-Bild und Textblöcken (Elemente verschieben)
d = pymupdf.open(); p = d.new_page(width=595, height=842)
p.insert_text((72, 90), 'Briefkopf Forstbetrieb', fontsize=20, fontname='hebo')
p.insert_text((72, 140), 'Dieser Absatz bleibt an seinem Platz.', fontsize=11, fontname='helv')
p.insert_text((72, 700), 'Fußnote: verschiebbarer Text', fontsize=10, fontname='helv')
p.insert_image(pymupdf.Rect(360, 200, 520, 320), filename=P('foto.jpg'))
d.save(P('bild.pdf'), garbage=3, deflate=True)
# Anmerkungen aus einem anderen Programm (Autor, Kommentar, Farben) zum Weiterbearbeiten
d = pymupdf.open(); p = d.new_page(width=595, height=842)
p.insert_text((72, 100), 'Vertragsentwurf zur Durchforstung 2026', fontsize=16, fontname='hebo')
p.insert_text((72, 140), 'Die Maßnahme beginnt im Herbst und endet vor der Brutzeit.', fontsize=11, fontname='helv')
def meta(a, who, txt=''):
    a.set_info(title=who, content=txt); a.update()
a = p.add_rect_annot(pymupdf.Rect(80, 200, 220, 280)); a.set_colors(stroke=(0.1, 0.3, 0.9)); a.set_border(width=2); meta(a, 'Anna Förster', 'Fläche prüfen')
a = p.add_circle_annot(pymupdf.Rect(300, 200, 400, 280)); a.set_colors(stroke=(0.9, 0.1, 0.1)); meta(a, 'Ben')
a = p.add_line_annot((80, 340), (260, 380)); a.set_line_ends(pymupdf.PDF_ANNOT_LE_NONE, pymupdf.PDF_ANNOT_LE_OPEN_ARROW); a.set_colors(stroke=(0, 0.6, 0.2)); meta(a, 'Ben')
a = p.add_freetext_annot(pymupdf.Rect(300, 330, 500, 370), 'Termin mit Revierleitung', fontsize=12, text_color=(0.8, 0, 0)); meta(a, 'Anna Förster', 'Termin mit Revierleitung')
q = p.search_for('Herbst')[0]; a = p.add_highlight_annot(q); a.set_colors(stroke=(1, 0.9, 0.2)); meta(a, 'Ben')
a = p.add_ink_annot([[(100, 450), (150, 500), (200, 450)]]); a.set_colors(stroke=(0.5, 0, 0.5)); a.set_border(width=3); meta(a, 'Ben')
a = p.add_text_annot((520, 90), 'Bitte gegenlesen', icon='Note'); meta(a, 'Anna Förster', 'Bitte gegenlesen')
d.save(P('kommentare.pdf'), garbage=3, deflate=True)
# 14 Alles, was beim Zusammenführen, Extrahieren und Teilen erhalten bleiben muss:
#    interne Links, benanntes Sprungziel, Webadresse, Formularfelder, Ebene, Anhang, Lesezeichen
d = text_pages(pymupdf.open(), 4, title='Reich')
p1 = d[0]
p1.insert_link({'kind': pymupdf.LINK_GOTO, 'from': pymupdf.Rect(72, 600, 300, 620), 'page': 2, 'to': pymupdf.Point(72, 90)})
p1.insert_link({'kind': pymupdf.LINK_URI, 'from': pymupdf.Rect(72, 630, 300, 650), 'uri': 'https://example.com/werkbank'})
p1.insert_text((72, 615), 'Zu Seite 3', fontsize=11); p1.insert_text((72, 645), 'Webseite', fontsize=11); p1.insert_text((72, 675), 'Zum Ziel', fontsize=11)
# benanntes Ziel „ziel“ auf Seite 4, Link darauf über eine GoTo-Aktion mit Namen
p4 = d[3].xref
names = d.get_new_xref(); d.update_object(names, f'<< /Names [(ziel) [{p4} 0 R /XYZ 72 750 0]] >>')
d.xref_set_key(d.pdf_catalog(), 'Names', f'<< /Dests {names} 0 R >>')
annot = d.get_new_xref()
d.update_object(annot, f'<< /Type /Annot /Subtype /Link /Rect [72 160 300 180] /Border [0 0 0] /A << /S /GoTo /D (ziel) >> >>')
d.xref_set_key(p1.xref, 'Annots', d.xref_get_key(p1.xref, 'Annots')[1].replace(']', f' {annot} 0 R]'))
# Formular auf Seite 2
p2 = d[1]
w = pymupdf.Widget(); w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT; w.field_name = 'Name'; w.rect = pymupdf.Rect(72, 600, 300, 622); w.field_value = 'Erika'; p2.add_widget(w)
w = pymupdf.Widget(); w.field_type = pymupdf.PDF_WIDGET_TYPE_CHECKBOX; w.field_name = 'Zustimmung'; w.rect = pymupdf.Rect(72, 640, 90, 658); w.field_value = True; p2.add_widget(w)
# Ebene auf Seite 3 (nur in der Ebene sichtbar)
oc = d.add_ocg('Ebene Massstab', on=True)
d[2].insert_text((72, 620), 'Nur in der Ebene', fontsize=14, oc=oc)
# Anhang und Lesezeichen
d.embfile_add('daten.csv', b'art;menge\nFichte;12\n', filename='daten.csv', desc='Quelldaten')
d.set_toc([[1, 'Links', 1], [1, 'Formular', 2], [1, 'Ebene', 3], [1, 'Ziel', 4]])
save(d, 'reich.pdf')

# 15 Root-Verweis ins Leere und falsche Querverweis-Position (abgebrochenes Schreiben): pdf.js lehnt ab, Reparatur rettet
_b = open(P('text.pdf'), 'rb').read().replace(b'/Root 1 0 R', b'/Root 99 0 R').replace(b'startxref\n6076', b'startxref\n9999')
open(P('kaputt-root.pdf'), 'wb').write(_b)

# Formular wie aus einer Textverarbeitung: eingebettete Teilschriften (Sans fett/normal, Serif, Mono) – braucht LibreOffice
import shutil, subprocess, tempfile
if shutil.which('soffice'):
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', tmp, P('src/formular.html')], check=True, capture_output=True)
        shutil.copy(os.path.join(tmp, 'formular.pdf'), P('formular.pdf'))
print('Korpus:', sorted(os.listdir(OUT)))
